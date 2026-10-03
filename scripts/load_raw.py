#!/usr/bin/env python3
"""Load the raw source files into the raw schema in PostgreSQL.

Creates the schemas and raw tables (sql/00 and sql/01), then loads:

    data/raw/ppd/pp-*.csv                Price Paid Data      -> raw.ppd
    data/raw/onspd/Data/ONSPD_*_UK.csv   Postcode Directory   -> raw.onspd
    data/raw/onspd/Documents/LAD ...csv  borough names        -> raw.lad_names
    data/raw/ashe/*Table 8.7a/8.7b*.xlsx earnings, with CVs   -> raw.ashe_table8
    data/raw/rents/*.xlsx, worksheet 2   rents by borough     -> raw.rents_borough

Values are loaded as text, unchanged. Each file's row count is compared with
the rows the database holds for it; the result goes to qa.check_log and a
mismatch stops the load with exit code 1.

Connection: --dsn, else DATABASE_URL, else
"host=localhost dbname=ppd_expansion user=postgres". Keep the password in
pgpass.conf (Windows) or ~/.pgpass, not in the command line.

    python scripts/load_raw.py
    python scripts/load_raw.py --dsn "host=localhost dbname=ppd_test user=postgres"
"""
import argparse
import csv
import datetime as dt
import hashlib
import os
import re
import sys
import time
from pathlib import Path

import openpyxl
import psycopg
from psycopg import sql

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
SQL_DIR = ROOT / "sql"
SETUP_SQL = ["00_create_schemas.sql", "01_raw_tables.sql"]
DEFAULT_DSN = "host=localhost dbname=ppd_expansion user=postgres"

PPD_YEARS = range(2015, 2026)  # brief: January 2015 to December 2025
META_COLUMNS = {"source_file", "row_num", "sheet"}

ASHE_HEADER = ["Description", "Code", "(thousand)", "Median", "change", "Mean", "change",
               "10", "20", "25", "30", "40", "60", "70", "75", "80", "90"]
RENTS_SHEET = "2"
RENTS_HEADER = ["Borough", "Bedroom Category", "Count of rents", "Mean",
                "Lower quartile", "Median", "Upper quartile"]


class CheckFailed(Exception):
    pass


class Loader:
    def __init__(self, conn, run_id):
        self.conn = conn
        self.run_id = run_id

    # ---- logging -------------------------------------------------------

    def check(self, name, subject, ok, expected, observed, detail=None):
        """Record a check; on failure commit the log, then stop."""
        status = "pass" if ok else "fail"
        self.conn.execute(
            "INSERT INTO qa.check_log (run_id, layer, check_name, subject, status, expected, observed, detail)"
            " VALUES (%s, 'raw', %s, %s, %s, %s, %s, %s)",
            (self.run_id, name, subject, status, str(expected), str(observed), detail),
        )
        if not ok:
            self.conn.commit()
            raise CheckFailed(f"{name} failed for {subject}: expected {expected}, observed {observed}"
                              + (f" ({detail})" if detail else ""))

    def log_load(self, table, path, file_bytes, sha, file_rows, loaded_rows):
        self.conn.execute(
            "INSERT INTO raw.load_log (run_id, table_name, source_file, file_bytes, file_sha256, file_rows, loaded_rows)"
            " VALUES (%s, %s, %s, %s, %s, %s, %s)",
            (self.run_id, table, rel(path), file_bytes, sha, file_rows, loaded_rows),
        )

    def reconcile(self, table, path, file_rows, sha):
        loaded = self.conn.execute(
            sql.SQL("SELECT count(*) FROM {} WHERE source_file = %s").format(sql.Identifier("raw", table)),
            (rel(path),),
        ).fetchone()[0]
        self.log_load(f"raw.{table}", path, path.stat().st_size, sha, file_rows, loaded)
        self.check("file_rows_equal_raw_rows", rel(path), loaded == file_rows, file_rows, loaded)
        self.conn.commit()
        print(f"  {rel(path)}: {loaded:,} rows")

    # ---- CSV sources ---------------------------------------------------

    def copy_csv(self, table, path, header):
        """Stream a CSV file into raw.<table> with COPY, tagging every row with its file."""
        target = sql.Identifier("raw", table)
        columns = [c for c in table_columns(self.conn, table) if c not in META_COLUMNS]
        if header:
            found = [h.strip().lower() for h in read_csv_header(path)]
            self.check("header_matches_table", rel(path), found == columns,
                       ",".join(columns), ",".join(found))
        # COPY cannot set a constant column, so source_file comes from a
        # column default that is set for the duration of this file's load.
        self.conn.execute(sql.SQL("ALTER TABLE {} ALTER COLUMN source_file SET DEFAULT {}")
                          .format(target, sql.Literal(rel(path))))
        stmt = sql.SQL("COPY {} ({}) FROM STDIN WITH (FORMAT csv, HEADER {}, ENCODING 'UTF8')").format(
            target, sql.SQL(", ").join(map(sql.Identifier, columns)),
            sql.SQL("true" if header else "false"))
        with self.conn.cursor() as cur, cur.copy(stmt) as copy, open(path, "rb") as f:
            while data := f.read(1 << 20):
                copy.write(data)
        self.conn.execute(sql.SQL("ALTER TABLE {} ALTER COLUMN source_file DROP DEFAULT").format(target))
        file_rows, sha = count_rows_and_hash(path, header)
        self.reconcile(table, path, file_rows, sha)

    def load_ppd(self):
        print("Price Paid Data -> raw.ppd")
        files = {p.name: p for p in (RAW / "ppd").glob("pp-*.csv")}
        expected = [f"pp-{y}.csv" for y in PPD_YEARS]
        missing = [f for f in expected if f not in files]
        self.check("expected_files_present", "data/raw/ppd", not missing,
                   f"{len(expected)} yearly files", f"missing {missing}" if missing else "all present")
        for name in expected:
            self.copy_csv("ppd", files[name], header=False)

    def load_onspd(self):
        print("ONS Postcode Directory -> raw.onspd, raw.lad_names")
        self.copy_csv("onspd", find_one(RAW / "onspd" / "Data", "ONSPD_*_UK.csv"), header=True)
        self.copy_csv("lad_names", find_one(RAW / "onspd" / "Documents",
                                            "LAD Local Authority District names and codes UK*.csv"), header=True)

    # ---- spreadsheet sources ---------------------------------------------

    def copy_rows(self, table, columns, rows):
        stmt = sql.SQL("COPY {} ({}) FROM STDIN").format(
            sql.Identifier("raw", table), sql.SQL(", ").join(map(sql.Identifier, columns)))
        with self.conn.cursor() as cur, cur.copy(stmt) as copy:
            for row in rows:
                copy.write_row(row)

    def load_ashe(self):
        print("ASHE Table 8.7a/8.7b -> raw.ashe_table8")
        files = sorted(p for p in (RAW / "ashe").glob("*.xlsx") if re.search(r"Table 8\.7[ab]\s", p.name))
        self.check("expected_files_present", "data/raw/ashe", len(files) == 2,
                   "Table 8.7a and 8.7b", [p.name for p in files])
        columns = table_columns(self.conn, "ashe_table8")
        for path in files:
            rows = []
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                for ws in wb.worksheets:
                    grid = list(ws.iter_rows(min_row=1, values_only=True))
                    head = find_header_row(grid, "Description")
                    if head is None:
                        continue  # the notes sheet
                    found = [header_text(v) for v in grid[head][:len(ASHE_HEADER)]]
                    self.check("header_matches_table", f"{rel(path)} [{ws.title}]",
                               found == ASHE_HEADER, ",".join(ASHE_HEADER), ",".join(map(str, found)))
                    for i, r in enumerate(grid[head + 1:], start=head + 2):
                        values = [cell_text(v) for v in r[:len(ASHE_HEADER)]]
                        if any(v is not None and v.strip() for v in values):
                            rows.append([rel(path), ws.title, i] + values)
            finally:
                wb.close()
            self.copy_rows("ashe_table8", columns, rows)
            self.reconcile("ashe_table8", path, len(rows), sha256(path))

    def load_rents(self):
        print("Private rents by borough -> raw.rents_borough")
        path = find_one(RAW / "rents", "*.xlsx")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            grid = list(wb[RENTS_SHEET].iter_rows(min_row=1, values_only=True))
        finally:
            wb.close()
        head = find_header_row(grid, "Borough")
        found = [header_text(v) for v in grid[head][:len(RENTS_HEADER)]] if head is not None else []
        self.check("header_matches_table", f"{rel(path)} [{RENTS_SHEET}]", found == RENTS_HEADER,
                   ",".join(RENTS_HEADER), ",".join(map(str, found)))
        rows = []
        for i, r in enumerate(grid[head + 1:], start=head + 2):
            values = [cell_text(v) for v in r[:len(RENTS_HEADER)]]
            if any(v is not None and v.strip() for v in values):
                rows.append([rel(path), i] + values)
        columns = table_columns(self.conn, "rents_borough")
        self.copy_rows("rents_borough", columns, rows)
        self.reconcile("rents_borough", path, len(rows), sha256(path))


# ---- helpers ---------------------------------------------------------------

def rel(path):
    return path.relative_to(ROOT).as_posix()


def find_one(folder, pattern):
    matches = sorted(folder.glob(pattern))
    if len(matches) != 1:
        raise CheckFailed(f"expected exactly one {pattern} in {folder}, found {len(matches)}")
    return matches[0]


def table_columns(conn, table):
    return [r[0] for r in conn.execute(
        "SELECT column_name FROM information_schema.columns"
        " WHERE table_schema = 'raw' AND table_name = %s ORDER BY ordinal_position", (table,))]


def read_csv_header(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return next(csv.reader(f))


def count_rows_and_hash(path, header):
    """Data rows (line count, minus the header) and SHA-256, in one pass.

    Counting lines matches a CSV row count only when no field contains a
    line break; true for these files (checked against a csv.reader count).
    """
    h, lines, last = hashlib.sha256(), 0, b"\n"
    with open(path, "rb") as f:
        while chunk := f.read(1 << 24):
            h.update(chunk)
            lines += chunk.count(b"\n")
            last = chunk[-1:]
    if last != b"\n":
        lines += 1  # final line has no line break
    return lines - (1 if header else 0), h.hexdigest()


def sha256(path):
    return count_rows_and_hash(path, False)[1]


def find_header_row(grid, first_cell):
    for i, r in enumerate(grid[:20]):
        if r and header_text(r[0]) == first_cell:
            return i
    return None


def cell_text(v):
    """Spreadsheet value as text, as it would read in the cell; None stays NULL."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def header_text(v):
    return (cell_text(v) or "").strip()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dsn", help="libpq connection string (default: DATABASE_URL or local ppd_expansion)")
    args = ap.parse_args()
    dsn = args.dsn or os.environ.get("DATABASE_URL") or DEFAULT_DSN

    run_id = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    started = time.monotonic()
    print(f"Raw load, run {run_id}")
    with psycopg.connect(dsn) as conn:
        for name in SETUP_SQL:
            conn.execute((SQL_DIR / name).read_text(encoding="utf-8"))
        conn.commit()
        loader = Loader(conn, run_id)
        try:
            loader.load_ppd()
            loader.load_onspd()
            loader.load_ashe()
            loader.load_rents()
        except CheckFailed as exc:
            conn.rollback()  # the failed check itself was already committed
            print(f"\nSTOPPED: {exc}", file=sys.stderr)
            return 1
    print(f"\nDone in {time.monotonic() - started:.0f}s. Results: qa.check_log and raw.load_log, run_id = '{run_id}'.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
